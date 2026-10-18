"""
Excel safety utilities shared across all apps.
"""
import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Font, PatternFill

_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def safe_cell(value):
    """
    Prevent CSV/Excel formula injection.
    Any cell value starting with a formula prefix character is escaped
    by prepending a single quote so Excel treats it as a literal string.
    """
    if value is None:
        return ""
    s = str(value)
    if s and s[0] in _FORMULA_PREFIXES:
        return "'" + s
    return s


def make_xlsx_response(filename):
    """Return an HttpResponse pre-configured for xlsx download."""
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def style_header_row(sheet):
    """Apply a dark-indigo header style to row 1 of the worksheet."""
    header_fill = PatternFill(start_color="3730A3", end_color="3730A3", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 20


def export_queryset_xlsx(queryset, columns, filename, sheet_title="Export"):
    """
    Build an .xlsx HttpResponse from a queryset.

    columns: list of (header, getter) tuples, where getter is a callable
    that takes one object from the queryset and returns its cell value
    (already run through safe_cell if it's free-text).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    headers = [h for h, _ in columns]
    ws.append(headers)
    style_header_row(ws)

    col_widths = {h: len(h) + 4 for h in headers}
    for obj in queryset:
        row = [getter(obj) for _, getter in columns]
        ws.append(row)
        for idx, val in enumerate(row):
            col_widths[headers[idx]] = max(col_widths[headers[idx]], len(str(val or "")) + 2)

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = col_widths[header]

    response = make_xlsx_response(filename)
    wb.save(response)
    return response
